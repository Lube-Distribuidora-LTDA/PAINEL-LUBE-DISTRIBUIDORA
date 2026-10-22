"""
LUBE OFFICE - gera os modelos em branco e os modelos com a marca da Lube.

Saida: office/modelos/*.docx, *.xlsx, *.pptx (servidos como arquivos estaticos
pelo portal; o navegador copia o modelo para o Storage ao criar um arquivo).

Rodar (na pasta do portal):  python office/ferramentas/gerar_modelos.py
Precisa de: python-docx 1.1, python-pptx 1.0, openpyxl 3.1, Pillow.
Cores e letras: nota "Identidade visual da Lube" no cerebro (azul #0D176E,
vermelho #E5101F; fonte Arial nos arquivos, que abre igual em qualquer maquina).
"""
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt as PptPt

RAIZ = Path(__file__).resolve().parents[1]
SAIDA = RAIZ / "modelos"
LOGO = RAIZ.parent / "assets" / "logo-lube.png"

AZUL = "0D176E"
VERMELHO = "E5101F"
AZUL_MEDIO = "2A3796"
AZUL_CLARO = "E8EBF7"
CINZA = "5C6480"
FONTE = "Arial"
ASSINATURA = "Lube Distribuidora · desde 1959 · Cariacica-ES"


def logo_leve(largura_px=640):
    """Logo reduzida (a original tem 1,2 MB) para nao pesar cada arquivo novo."""
    img = Image.open(LOGO).convert("RGBA")
    alt = round(img.height * largura_px / img.width)
    img = img.resize((largura_px, alt), Image.LANCZOS)
    buf = BytesIO()
    img.save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


# Acabamento de todo arquivo gerado (achado no teste de ponta a ponta, 2026-10-21):
# as bibliotecas deixam o idioma em ingles (o corretor sublinhava tudo) e as
# propriedades com "python-docx", data de 2013 e "Microsoft Macintosh Word".
def finalizar(caminho, titulo=""):
    agora = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    core = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f"<dc:title>{escape(titulo)}</dc:title><dc:creator>Lube Distribuidora</dc:creator>"
        "<cp:lastModifiedBy>Lube Distribuidora</cp:lastModifiedBy><cp:revision>1</cp:revision>"
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{agora}</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{agora}</dcterms:modified></cp:coreProperties>'
    )
    tmp = caminho.with_name(caminho.name + ".tmp")
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename.endswith(".xml"):
                txt = dados.decode("utf-8")
                if item.filename == "docProps/core.xml":
                    txt = core
                elif item.filename == "docProps/app.xml":
                    txt = re.sub(r"<Application>.*?</Application>", "<Application>LUBE OFFICE</Application>", txt)
                    txt = re.sub(r"<AppVersion>.*?</AppVersion>", "", txt)
                txt = (txt.replace('w:val="en-US"', 'w:val="pt-BR"').replace('lang="en-US"', 'lang="pt-BR"')
                          .replace('w:eastAsia="ja-JP"', 'w:eastAsia="pt-BR"').replace('w:eastAsia="en-US"', 'w:eastAsia="pt-BR"'))
                dados = txt.encode("utf-8")
            zout.writestr(item, dados)
    tmp.replace(caminho)


# ---------------------------------------------------------------- Word
def doc_base():
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.0)
    normal = d.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for nome, tam, cor in (("Title", 26, AZUL), ("Heading 1", 15, AZUL), ("Heading 2", 12, AZUL_MEDIO)):
        st = d.styles[nome]
        st.font.name = FONTE
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(cor)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
        st.paragraph_format.space_before = Pt(14 if nome != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)
    return d


def faixa_vermelha(par, largura_cm=4.0):
    """Filete vermelho embaixo do titulo (borda inferior do paragrafo)."""
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "18")
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), VERMELHO)
    bdr.append(b)
    pPr.append(bdr)


def cabecalho_rodape(d, com_logo=True):
    sec = d.sections[0]
    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if com_logo:
        cab.add_run().add_picture(logo_leve(), width=Cm(3.6))
    rod = sec.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = rod.add_run(ASSINATURA)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(CINZA)


def sombrear(celula, cor):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def tabela(d, cabecalho, linhas, larguras=None):
    t = d.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(cabecalho):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(txt)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        sombrear(c, AZUL)
    for lin in linhas:
        cells = t.add_row().cells
        for i, txt in enumerate(lin):
            cells[i].text = txt
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t


def campos(d, pares):
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for rot, val in pares:
        cells = t.add_row().cells
        cells[0].text = ""
        r = cells[0].paragraphs[0].add_run(rot)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(AZUL)
        sombrear(cells[0], AZUL_CLARO)
        cells[1].text = val
        cells[0].width, cells[1].width = Cm(4.5), Cm(12.1)
    d.add_paragraph()


def ajuda(d, texto):
    p = d.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(CINZA)
    return p


def branco_docx():
    d = doc_base()
    d.add_paragraph("")
    d.save(SAIDA / "branco.docx")


def proposta_docx():
    d = doc_base()
    cabecalho_rodape(d)
    t = d.add_paragraph("Proposta comercial", style="Title")
    faixa_vermelha(t)
    campos(d, [("Cliente", "Nome do cliente"), ("CNPJ", "00.000.000/0000-00"), ("Contato", "Nome · telefone · e-mail"),
               ("Data", "AAAA-MM-DD"), ("Validade", "15 dias"), ("RCA responsável", "Nome do RCA")])
    d.add_paragraph("1. Apresentação", style="Heading 1")
    d.add_paragraph("A Lube Distribuidora atende o Espírito Santo desde 1959 com distribuição de alimentos, bebidas, "
                    "higiene e limpeza. Esta proposta reúne os produtos, as condições e o prazo combinados com o cliente.")
    d.add_paragraph("2. Produtos e preços", style="Heading 1")
    tabela(d, ["Código", "Produto", "Qtde", "Preço unit. (R$)", "Total (R$)"],
           [["", "", "", "", ""] for _ in range(6)], [2.2, 7.0, 1.8, 2.8, 2.8])
    p = d.add_paragraph()
    r = p.add_run("Total da proposta: R$ ")
    r.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph("3. Condições comerciais", style="Heading 1")
    for txt in ("Forma de pagamento: ", "Prazo de pagamento: ", "Pedido mínimo: ", "Frete: "):
        d.add_paragraph(txt, style="List Bullet")
    d.add_paragraph("4. Entrega", style="Heading 1")
    d.add_paragraph("Prazo de entrega a partir da confirmação do pedido: ")
    d.add_paragraph("5. Aceite", style="Heading 1")
    ajuda(d, "Assinar e devolver ao RCA. Os preços valem até a data de validade.")
    t = d.add_table(rows=2, cols=2)
    for i, quem in enumerate(("Lube Distribuidora", "Cliente")):
        t.rows[0].cells[i].text = "\n\n______________________________"
        t.rows[1].cells[i].text = quem
        t.rows[1].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(SAIDA / "proposta-comercial.docx")


def ata_docx():
    d = doc_base()
    cabecalho_rodape(d)
    t = d.add_paragraph("Ata de reunião", style="Title")
    faixa_vermelha(t)
    campos(d, [("Assunto", ""), ("Data e hora", "AAAA-MM-DD, 00:00 às 00:00"), ("Local", "Sala de reunião · Teams"),
               ("Conduzida por", ""), ("Registrada por", "")])
    d.add_paragraph("Participantes", style="Heading 1")
    for _ in range(3):
        d.add_paragraph("", style="List Bullet")
    d.add_paragraph("Pauta", style="Heading 1")
    for _ in range(3):
        d.add_paragraph("", style="List Number")
    d.add_paragraph("O que foi discutido", style="Heading 1")
    ajuda(d, "Resumo de cada ponto da pauta, com os números e os motivos.")
    d.add_paragraph("Decisões e responsáveis", style="Heading 1")
    tabela(d, ["Decisão", "Responsável", "Prazo", "Situação"], [["", "", "", "A fazer"] for _ in range(5)], [7.6, 3.6, 2.4, 3.0])
    d.add_paragraph("Próxima reunião", style="Heading 1")
    d.add_paragraph("Data: ")
    d.save(SAIDA / "ata-de-reuniao.docx")


# ---------------------------------------------------------------- Excel
def branco_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    wb.save(SAIDA / "branco.xlsx")


def relatorio_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Indicadores"
    ws.sheet_view.showGridLines = False
    azul = PatternFill("solid", fgColor=AZUL)
    claro = PatternFill("solid", fgColor=AZUL_CLARO)
    fino = Side(style="thin", color="D5DAEA")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    larguras = {"A": 3, "B": 34, "C": 16, "D": 14, "E": 14, "F": 12, "G": 16, "H": 40}
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 44
    img = XlImage(logo_leve(360))
    img.width, img.height = 150, 50
    ws.add_image(img, "G2")
    ws["B2"] = "Relatório mensal"
    ws["B2"].font = Font(name=FONTE, size=22, bold=True, color=AZUL)
    ws["B3"] = "Mês de referência:"
    ws["B3"].font = Font(name=FONTE, size=10, color=CINZA)
    ws["C3"] = "AAAA-MM"
    ws["C3"].font = Font(name=FONTE, size=10, bold=True, color=AZUL)
    for col in "BCDEFGH":
        ws[f"{col}4"].fill = PatternFill("solid", fgColor=VERMELHO)
    ws.row_dimensions[4].height = 4

    cab = ["Indicador", "Área", "Unidade", "Meta", "Realizado", "% da meta", "Comentário"]
    for i, txt in enumerate(cab):
        c = ws.cell(row=6, column=2 + i, value=txt)
        c.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(vertical="center", horizontal="left" if i in (0, 1, 6) else "center")
        c.border = borda
    ws.row_dimensions[6].height = 24
    exemplos = [("Venda", "Comercial", "R$"), ("Margem", "Comercial", "%"), ("Ruptura", "Compras", "%"),
                ("Devoluções", "Comercial", "R$"), ("Entregas no prazo", "Logística", "%"),
                ("Absenteísmo", "RH", "%"), ("Contas pagas em dia", "Financeiro", "%"), ("", "", ""), ("", "", ""), ("", "", "")]
    for n, (ind, area, un) in enumerate(exemplos):
        r = 7 + n
        valores = [ind, area, un, None, None, f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r}),E{r}<>0),F{r}/E{r},"")', None]
        for i, v in enumerate(valores):
            c = ws.cell(row=r, column=2 + i, value=v)
            c.font = Font(name=FONTE, size=10, color="1B2140")
            c.border = borda
            c.alignment = Alignment(vertical="center", horizontal="left" if i in (0, 1, 6) else "center")
            if n % 2:
                c.fill = claro
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        ws.cell(row=r, column=6).number_format = "#,##0.00"
        ws.cell(row=r, column=7).number_format = "0.0%"
        ws.row_dimensions[r].height = 20
    ultima = 7 + len(exemplos) - 1
    alvo = f"G7:G{ultima}"
    ws.conditional_formatting.add(alvo, CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                                   font=Font(color="15965A", bold=True)))
    ws.conditional_formatting.add(alvo, CellIsRule(operator="lessThan", formula=["0.9"],
                                                   font=Font(color=VERMELHO, bold=True)))
    dv = DataValidation(type="list", formula1='"Comercial,Compras,Financeiro,Fiscal,Logística,RH,TI"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C7:C{ultima}")
    nota = ws.cell(row=ultima + 2, column=2,
                   value="Verde: meta batida. Vermelho: abaixo de 90% da meta. Para indicador em que menos é melhor (ruptura, "
                         "devoluções, absenteísmo), escreva a leitura no comentário.")
    nota.font = Font(name=FONTE, size=9, italic=True, color=CINZA)
    rod = ws.cell(row=ultima + 3, column=2, value=ASSINATURA)
    rod.font = Font(name=FONTE, size=8, color=CINZA)
    ws.freeze_panes = "B7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SAIDA / "relatorio-mensal.xlsx")


# ---------------------------------------------------------------- PowerPoint
def rgb(h):
    return PptRGB.from_string(h)


def texto(slide, x, y, w, h, txt, tam, cor="FFFFFF", negrito=False, alinhar=PP_ALIGN.LEFT):
    caixa = slide.shapes.add_textbox(x, y, w, h)
    tf = caixa.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alinhar
    r = p.add_run()
    r.text = txt
    r.font.size = PptPt(tam)
    r.font.bold = negrito
    r.font.name = FONTE
    r.font.language_id = MSO_LANGUAGE_ID.BRAZILIAN_PORTUGUESE
    r.font.color.rgb = rgb(cor)
    return caixa


def faixa(slide, x, y, w, h, cor):
    f = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, x, y, w, h)
    f.fill.solid()
    f.fill.fore_color.rgb = rgb(cor)
    f.line.fill.background()
    f.adjustments[0] = 0.25
    return f


def placa_logo(slide, x, y, largura):
    altura = Emu(int(largura * 0.36))
    placa = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, largura, altura)
    placa.fill.solid()
    placa.fill.fore_color.rgb = rgb("FFFFFF")
    placa.line.fill.background()
    placa.adjustments[0] = 0.12
    margem = int(largura * 0.08)
    slide.shapes.add_picture(logo_leve(), x + margem, y + int(altura * 0.12), width=largura - 2 * margem)


def fundo(slide, cor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(cor)


def apresentacao(com_conteudo):
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    vazio = p.slide_layouts[6]
    W, H = p.slide_width, p.slide_height
    if not com_conteudo:
        s = p.slides.add_slide(vazio)
        texto(s, Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.2), "Título da apresentação", 40, "1B2140", True)
        p.save(SAIDA / "branco.pptx")
        return

    # capa
    s = p.slides.add_slide(vazio)
    fundo(s, AZUL)
    faixa(s, Inches(-0.6), Inches(5.2), Inches(8.2), Inches(0.55), VERMELHO)
    placa_logo(s, Inches(10.0), Inches(0.6), Inches(2.7))
    texto(s, Inches(0.8), Inches(2.3), Inches(10), Inches(1.4), "Título da apresentação", 44, "FFFFFF", True)
    texto(s, Inches(0.8), Inches(3.6), Inches(10), Inches(0.8), "Área · mês de referência", 20, "C9D0F0")
    texto(s, Inches(0.8), Inches(6.5), Inches(10), Inches(0.5), ASSINATURA, 11, "8E99D6")

    def slide_titulo(titulo):
        sl = p.slides.add_slide(vazio)
        fundo(sl, "FFFFFF")
        faixa(sl, Inches(0.6), Inches(0.62), Inches(0.5), Inches(0.5), VERMELHO)
        texto(sl, Inches(1.25), Inches(0.45), Inches(9.5), Inches(0.9), titulo, 30, AZUL, True)
        placa = sl.shapes.add_picture(logo_leve(), Inches(11.1), Inches(0.5), width=Inches(1.7))
        rodape = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.12), W, Inches(0.12))
        rodape.fill.solid()
        rodape.fill.fore_color.rgb = rgb(AZUL)
        rodape.line.fill.background()
        return sl

    # agenda
    s = slide_titulo("Agenda")
    for i, item in enumerate(("Resultado do mês", "O que funcionou e o que não funcionou", "Próximos passos")):
        y = Inches(1.9 + i * 1.3)
        n = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.25), y, Inches(0.8), Inches(0.8))
        n.fill.solid()
        n.fill.fore_color.rgb = rgb(AZUL)
        n.line.fill.background()
        n.text_frame.text = str(i + 1)
        r = n.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = PptPt(20), True, FONTE
        r.font.color.rgb = rgb("FFFFFF")
        n.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        texto(s, Inches(2.4), y + Inches(0.08), Inches(9), Inches(0.7), item, 24, "1B2140")

    # resultado em números
    s = slide_titulo("Resultado do mês")
    for i, (rot, val) in enumerate((("Venda", "R$ 0,0 mi"), ("Margem", "0,0%"), ("Meta atingida", "0%"))):
        x = Inches(1.25 + i * 3.85)
        caixa = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(3.5), Inches(2.6))
        caixa.fill.solid()
        caixa.fill.fore_color.rgb = rgb(AZUL_CLARO)
        caixa.line.fill.background()
        caixa.adjustments[0] = 0.08
        texto(s, x + Inches(0.3), Inches(2.3), Inches(2.9), Inches(0.6), rot.upper(), 14, AZUL_MEDIO, True)
        texto(s, x + Inches(0.3), Inches(3.0), Inches(2.9), Inches(1.0), val, 40, AZUL, True)
    texto(s, Inches(1.25), Inches(5.2), Inches(11), Inches(1.0),
          "Uma frase que explica o resultado: o que puxou para cima, o que puxou para baixo.", 18, "5C6480")

    # próximos passos
    s = slide_titulo("Próximos passos")
    caixa = s.shapes.add_textbox(Inches(1.25), Inches(1.8), Inches(11), Inches(4.5))
    tf = caixa.text_frame
    tf.word_wrap = True
    for i, item in enumerate(("Ação · responsável · prazo", "Ação · responsável · prazo", "Ação · responsável · prazo")):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = par.add_run()
        r.text = "▸  " + item
        r.font.size, r.font.name = PptPt(24), FONTE
        r.font.language_id = MSO_LANGUAGE_ID.BRAZILIAN_PORTUGUESE
        r.font.color.rgb = rgb("1B2140")
        par.space_after = PptPt(18)

    # encerramento
    s = p.slides.add_slide(vazio)
    fundo(s, AZUL)
    faixa(s, Inches(5.2), Inches(4.35), Inches(8.6), Inches(0.5), VERMELHO)
    placa_logo(s, Inches(5.0), Inches(1.7), Inches(3.3))
    texto(s, Inches(0.8), Inches(5.3), Inches(11.7), Inches(0.9), "Obrigado", 36, "FFFFFF", True, PP_ALIGN.CENTER)
    p.save(SAIDA / "apresentacao-lube.pptx")


# ---------------------------------------------------------------- IA da Lube (2026-10-21)
# Bases dos arquivos que a IA cria pelo Drive: o plugin office/ia/executor.js escreve o
# conteúdo dentro do editor. O documento já nasce com o cabeçalho e o rodapé da Lube; a
# apresentação, em 16:9 com um slide vazio (o plugin troca pelos slides da IA).
def ia_documento():
    d = doc_base()
    cabecalho_rodape(d)
    d.add_paragraph("")
    d.save(SAIDA / "ia-documento.docx")


def ia_apresentacao():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    p.slides.add_slide(p.slide_layouts[6])
    p.save(SAIDA / "ia-apresentacao.pptx")


def logo_plugin():
    """Logo que o plugin põe nos slides (o motor baixa do portal) e o ícone do plugin."""
    pasta = RAIZ / "ia"
    pasta.mkdir(exist_ok=True)
    (pasta / "logo-lube.png").write_bytes(logo_leve(480).getvalue())   # vai em cada slide: leve
    simb = Image.open(LOGO).convert("RGBA")
    simb = simb.crop((0, 0, round(simb.width * 0.37), simb.height))   # o "L" azul com o quadrado vermelho
    for nome, lado in (("icone.png", 40), ("icone@2x.png", 80)):
        quadro = Image.new("RGBA", (lado, lado), (0, 0, 0, 0))
        s = simb.copy()
        s.thumbnail((lado, lado), Image.LANCZOS)
        quadro.paste(s, ((lado - s.width) // 2, (lado - s.height) // 2), s)
        quadro.save(pasta / nome, "PNG", optimize=True)


if __name__ == "__main__":
    import sys
    SAIDA.mkdir(parents=True, exist_ok=True)
    if sys.argv[1:] == ["ia"]:
        # só as bases da IA, sem regravar os outros modelos
        ia_documento()
        ia_apresentacao()
        logo_plugin()
        finalizar(SAIDA / "ia-documento.docx", "")
        finalizar(SAIDA / "ia-apresentacao.pptx", "")
        sys.exit(0)
    branco_docx()
    branco_xlsx()
    apresentacao(False)
    proposta_docx()
    ata_docx()
    relatorio_xlsx()
    apresentacao(True)
    ia_documento()
    ia_apresentacao()
    logo_plugin()
    titulos = {"proposta-comercial.docx": "Proposta comercial", "ata-de-reuniao.docx": "Ata de reunião",
               "relatorio-mensal.xlsx": "Relatório mensal", "apresentacao-lube.pptx": "Apresentação Lube"}
    for f in sorted(SAIDA.iterdir()):
        if f.suffix in (".docx", ".xlsx", ".pptx"):
            finalizar(f, titulos.get(f.name, ""))
    for f in sorted(SAIDA.iterdir()):
        print(f"{f.name:28} {f.stat().st_size / 1024:7.1f} KB")
